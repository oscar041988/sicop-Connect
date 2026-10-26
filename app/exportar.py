"""Exportar a Excel el listado filtrado del panel (A5).

* Respeta los filtros activos: usa la misma consulta que la lista (panel.listar), sin paginar.
* Identificadores como TEXTO: Número SICOP y N.º de procedimiento se guardan como cadena y con formato «@», así
  Excel no los convierte en número (perdería ceros iniciales o los mostraría como 2,02608E+10).
* Ningún texto se guarda como fórmula: un título que empiece con «=» quedaría ejecutable en Excel (inyección de
  fórmulas). Todo texto se fuerza a celda de texto.
* Montos: una columna por moneda (Decimal, nunca convertidos ni sumados entre monedas).
* Una hoja «Filtros» dice qué filtros se aplicaron, cuándo se generó y de dónde salen los datos.
Dependencia: openpyxl (justificada en docs/AMPLIACIONES.md, A5).
"""

from __future__ import annotations

import io
import sqlite3
from datetime import datetime

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app import panel
from app.importers.comun import HORA_CR

TEXTO = "@"
FORMATO_FECHA = "dd/mm/yyyy hh:mm"
FORMATO_MONTO = "#,##0.00"
ENCABEZADO = PatternFill("solid", fgColor="E3EAF2")


def _texto(celda: Cell, valor: str | None) -> None:
    """Celda de texto: nunca número ni fórmula. Vacía si no hay dato, pero con formato texto igual."""
    celda.number_format = TEXTO
    if valor is None or valor == "":
        return
    celda.value = str(valor)
    celda.data_type = "s"


def _fecha(iso: str | None) -> datetime | None:
    """Fecha de SICOP → fecha y hora de Costa Rica sin zona (Excel no guarda zonas horarias)."""
    if not iso:
        return None
    try:
        fecha = datetime.fromisoformat(iso)
    except ValueError:
        return None
    return fecha.astimezone(HORA_CR).replace(tzinfo=None) if fecha.tzinfo else fecha


def descripcion_filtros(conn: sqlite3.Connection, f: panel.Filtros) -> list[tuple[str, str]]:
    """Los filtros como se ven en el panel, para la hoja «Filtros»."""
    prioridad = {"candidatos": "Candidatos (Alta y Media)", "todas": "Todas"}.get(
        f.prioridad, panel.ENCAJES.get(f.prioridad, f.prioridad))
    institucion = "Todas"
    if f.institucion:
        fila = conn.execute("SELECT nombre FROM institutions WHERE cedula = ?", (f.institucion,)).fetchone()
        institucion = f"{fila[0] if fila and fila[0] else 'Sin nombre'} (cédula {f.institucion})"
    revisado = {"": "Todos", "si": "Revisados", "no": "Sin revisar"}.get(f.revisado) or panel.DECISIONES.get(f.revisado, "")
    return [
        ("Prioridad", prioridad),
        ("Familia", panel.NOMBRES_FAMILIA.get(f.familia, f.familia) if f.familia else "Todas"),
        ("Institución", institucion),
        ("Plazo", panel.PLAZOS.get(f.plazo, f.plazo) if f.plazo else "Vigentes (sin vencidos)"),
        ("Revisión", revisado),
        ("Tipo de procedimiento", f.tipo or "Todos"),
        ("Búsqueda", f.q or "—"),
    ]


COLUMNAS = [  # (título, ancho)
    ("Número SICOP", 15), ("N.º de procedimiento", 28), ("Objeto del cartel", 60), ("Institución", 40),
    ("Familia", 24), ("Servicio de Hermatix", 18), ("Prioridad", 11), ("Plazo", 22), ("Cierre de ofertas", 17),
    ("Tipo de procedimiento", 26), ("Revisión", 18), ("Nuevo", 8), ("Palabras clave", 40),
]


def libro(conn: sqlite3.Connection, version: str, filtros: panel.Filtros, ahora: datetime,
          base_url: str) -> tuple[bytes, int]:
    """(contenido del .xlsx, filas exportadas) con todas las oportunidades que muestra el panel con esos filtros."""
    f = filtros.limpios()
    pagina = panel.listar(conn, version, f, ahora, todas=True)
    monedas = sorted({m for fila in pagina.filas for m in fila.estimado})

    wb = Workbook()
    hoja = wb.active
    hoja.title = "Oportunidades"
    titulos = [t for t, _ in COLUMNAS] + [f"Monto estimado {m}" for m in monedas] + ["Enlace al panel"]
    anchos = [a for _, a in COLUMNAS] + [18] * len(monedas) + [45]
    for col, (titulo, ancho) in enumerate(zip(titulos, anchos), start=1):
        celda = hoja.cell(row=1, column=col)
        _texto(celda, titulo)
        celda.font, celda.fill = Font(bold=True), ENCABEZADO
        hoja.column_dimensions[get_column_letter(col)].width = ancho

    for n, fila in enumerate(pagina.filas, start=2):
        # legible: como en el panel (SICOP usa espacios especiales que se ven raros en Excel).
        textos = [fila.nro_sicop, fila.nro_procedimiento, panel.legible(panel.formato_titulo(fila.descripcion or "")),
                  panel.legible(fila.institucion), panel.NOMBRES_FAMILIA.get(fila.familia, fila.familia), fila.producto,
                  panel.ENCAJES.get(fila.encaje, fila.encaje), panel.PLAZOS.get(fila.plazo, fila.plazo)]
        for col, valor in enumerate(textos, start=1):
            _texto(hoja.cell(row=n, column=col), valor)
        cierre = hoja.cell(row=n, column=9)
        cierre.value = _fecha(fila.fecha_cierre)
        cierre.number_format = FORMATO_FECHA
        for col, valor in enumerate([fila.tipo_procedimiento,
                                     panel.DECISIONES.get(fila.decision, "Sin revisar") if fila.decision else "Sin revisar",
                                     "Sí" if fila.es_nuevo else "No", panel.legible(", ".join(fila.frases))], start=10):
            _texto(hoja.cell(row=n, column=col), valor)
        for col, moneda in enumerate(monedas, start=len(COLUMNAS) + 1):
            monto = fila.estimado.get(moneda)
            if monto is not None:
                celda = hoja.cell(row=n, column=col, value=monto.valor)  # Decimal: exacto, sin float
                celda.number_format = FORMATO_MONTO
        enlace = hoja.cell(row=n, column=len(titulos))
        _texto(enlace, f"{base_url.rstrip('/')}/oportunidades/{fila.nro_sicop}")
        enlace.hyperlink = enlace.value
        hoja.cell(row=n, column=3).alignment = Alignment(wrap_text=False)
    hoja.freeze_panes = "B2"
    hoja.auto_filter.ref = f"A1:{get_column_letter(len(titulos))}{max(len(pagina.filas) + 1, 1)}"

    info = wb.create_sheet("Filtros")
    filas_info = [("Exportado", ahora.astimezone(HORA_CR).strftime("%d/%m/%Y %H:%M") + " (hora de Costa Rica)"),
                  ("Oportunidades", str(len(pagina.filas))),
                  *descripcion_filtros(conn, f),
                  ("Reglas", version),
                  ("Fuente", "Datos abiertos de SICOP. La prioridad indica encaje con el portafolio de Hermatix, "
                             "no que se pueda participar. Un cartel aparece una vez por cada familia con la que coincide."),
                  ("Montos", "Estimado = cantidad × precio unitario de las líneas, por moneda; nunca se convierten.")]
    for n, (clave, valor) in enumerate(filas_info, start=1):
        _texto(info.cell(row=n, column=1), clave)
        info.cell(row=n, column=1).font = Font(bold=True)
        _texto(info.cell(row=n, column=2), valor)
    info.column_dimensions["A"].width = 24
    info.column_dimensions["B"].width = 100

    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue(), len(pagina.filas)
