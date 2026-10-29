"""«Ver» un documento guardado de SICOP sin descargarlo (pedido de Oscar, 2026-10-29).

* PDF e imágenes: el navegador los muestra solo; la ruta los manda «en línea» (inline) con su tipo fijo.
* ZIP: la lista de lo que trae; cada PDF, imagen o Excel de adentro se puede ver igual (no se descomprime a disco).
* Excel (.xlsx, .xlsm): cada hoja como tabla, con openpyxl (ya es dependencia por la exportación). Los valores se
  muestran como texto; las macros de un .xlsm nunca se ejecutan.
Word y lo demás: solo «Descargar». Solo lee: no escribe en la base ni en el disco.
"""

from __future__ import annotations

import io
import logging
import warnings
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path, PurePosixPath

log = logging.getLogger(__name__)

# El tipo de vista sale de la extensión. SVG queda fuera a propósito: puede traer código.
VISTAS = {".pdf": "pdf", ".jpg": "imagen", ".jpeg": "imagen", ".png": "imagen", ".zip": "zip",
          ".xlsx": "excel", ".xlsm": "excel"}
TIPOS = {".pdf": "application/pdf", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png"}
MAX_ENTRADAS_ZIP = 200   # igual que al leer el texto (sicop_documentos; no se importa: cargaría Playwright)
MAX_FILAS = 500       # por hoja: una lista de insumos larga se ve completa al descargarla
MAX_COLUMNAS = 50


def _extension(nombre: str | None) -> str:
    return PurePosixPath((nombre or "").replace("\\", "/")).suffix.lower()


def vista(nombre: str | None) -> str | None:
    """'pdf', 'imagen', 'zip', 'excel' o None (solo se puede descargar)."""
    return VISTAS.get(_extension(nombre))


def tipo(nombre: str | None) -> str | None:
    """El tipo con que se manda un PDF o una imagen en línea; None para lo demás."""
    return TIPOS.get(_extension(nombre))


@dataclass
class Entrada:
    indice: int
    nombre: str
    tamano: int
    vista: str | None


@dataclass
class Hoja:
    nombre: str
    filas: list[list[str]] = field(default_factory=list)
    recortada: bool = False


class NoSePuedeLeer(ValueError):
    """El archivo no se pudo abrir para verlo (dañado, con contraseña, demasiado grande…)."""


def entradas_zip(ruta: Path) -> list[Entrada]:
    """Los archivos de un ZIP (sin carpetas), en su orden. El índice es el que usa leer_entrada()."""
    try:
        with zipfile.ZipFile(ruta) as zf:
            infos = zf.infolist()
    except (zipfile.BadZipFile, OSError) as exc:
        raise NoSePuedeLeer("el ZIP no se puede abrir; descárguelo y ábralo en la computadora") from exc
    if len(infos) > MAX_ENTRADAS_ZIP:
        raise NoSePuedeLeer(f"el ZIP trae {len(infos)} archivos (máximo {MAX_ENTRADAS_ZIP}); descárguelo")
    return [Entrada(i, info.filename, info.file_size, vista(info.filename))
            for i, info in enumerate(infos) if not info.is_dir()]


def leer_entrada(ruta: Path, indice: int, max_bytes: int) -> tuple[str, bytes] | None:
    """Nombre y contenido de una entrada del ZIP, o None si no existe. Nunca lee más de max_bytes, aunque el ZIP
    mienta sobre el tamaño (un ZIP pequeño puede descomprimirse en gigas)."""
    try:
        with zipfile.ZipFile(ruta) as zf:
            infos = zf.infolist()
            if not 0 <= indice < len(infos) or infos[indice].is_dir():
                return None
            info = infos[indice]
            if info.file_size > max_bytes:
                raise NoSePuedeLeer(f"descomprimido pesa {info.file_size / 1_000_000:.1f} MB; descargue el ZIP")
            with zf.open(info) as f:
                datos = f.read(max_bytes + 1)
    except (zipfile.BadZipFile, OSError, RuntimeError) as exc:   # RuntimeError: ZIP con contraseña
        raise NoSePuedeLeer("no se pudo abrir ese archivo del ZIP") from exc
    if len(datos) > max_bytes:
        raise NoSePuedeLeer("descomprimido pasa del máximo; descargue el ZIP")
    return info.filename, datos


def _celda(valor) -> str:
    if valor is None:
        return ""
    if isinstance(valor, datetime):
        return valor.strftime("%d/%m/%Y %H:%M") if valor.time() != time(0) else valor.strftime("%d/%m/%Y")
    if isinstance(valor, date):
        return valor.strftime("%d/%m/%Y")
    if isinstance(valor, bool):
        return "Sí" if valor else "No"
    return str(valor)


def hojas_excel(datos: bytes) -> list[Hoja]:
    """Cada hoja con sus valores (lo calculado, no las fórmulas), sin las filas y columnas vacías del final."""
    with warnings.catch_warnings():   # «extension is not supported»: validaciones y formatos que no se ven
        warnings.simplefilter("ignore", UserWarning)
        return _hojas_excel(datos)


def _hojas_excel(datos: bytes) -> list[Hoja]:
    from openpyxl import load_workbook   # se carga solo al ver un Excel
    try:
        libro = load_workbook(io.BytesIO(datos), read_only=True, data_only=True)
    except Exception as exc:   # openpyxl lanza varios tipos según el daño del archivo
        log.info("Excel que no se pudo abrir para verlo: %s", type(exc).__name__)
        raise NoSePuedeLeer("el Excel está dañado, tiene contraseña o no es un .xlsx") from exc
    try:
        hojas = []
        for hoja in libro.worksheets:
            filas, recortada = [], False
            for n, fila in enumerate(hoja.iter_rows(max_col=MAX_COLUMNAS, values_only=True)):
                if n == MAX_FILAS:
                    recortada = True
                    break
                filas.append([_celda(v) for v in fila])
            while filas and not any(filas[-1]):
                filas.pop()
            ancho = max((max((i + 1 for i, v in enumerate(f) if v), default=0) for f in filas), default=0)
            hojas.append(Hoja(hoja.title, [f[:ancho] + [""] * (ancho - len(f)) for f in filas], recortada))
        return hojas
    except Exception as exc:
        log.info("Excel que no se pudo leer para verlo: %s", type(exc).__name__)
        raise NoSePuedeLeer("no se pudo leer el Excel; descárguelo") from exc
    finally:
        libro.close()
