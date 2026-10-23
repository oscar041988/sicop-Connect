"""A5: «Exportar a Excel» respeta los filtros activos y deja los identificadores como texto.

Datos SINTÉTICOS (los carteles de tests/test_panel.py) en una base temporal; el .xlsx se lee con openpyxl.
"""

import io
from contextlib import closing
from decimal import Decimal

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.database import connect
from app.main import app
from tests.test_panel import cliente, db  # noqa: F401

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _hoja(respuesta):
    assert respuesta.status_code == 200 and respuesta.headers["content-type"] == XLSX
    assert respuesta.headers["content-disposition"].startswith('attachment; filename="oportunidades_')
    libro = load_workbook(io.BytesIO(respuesta.content))
    hoja = libro["Oportunidades"]
    titulos = [c.value for c in hoja[1]]
    filas = [dict(zip(titulos, (c for c in fila))) for fila in hoja.iter_rows(min_row=2)]
    return libro, titulos, filas


def test_exporta_con_los_filtros_por_defecto(cliente):
    _, titulos, filas = _hoja(cliente.get("/oportunidades.xlsx"))
    assert titulos[:3] == ["Número SICOP", "N.º de procedimiento", "Objeto del cartel"]
    # Candidatos (Alta y Media), como la lista por defecto.
    assert sorted(f["Número SICOP"].value for f in filas) == ["00000000001", "00000000002", "00000000004"]


def test_los_ids_quedan_como_texto(cliente, db):
    with closing(connect(db)) as conn:
        conn.execute("UPDATE tenders SET nro_procedimiento = '0002026LD-000214-0000900001' WHERE nro_sicop = '00000000001'")
        conn.commit()
    _, _, filas = _hoja(cliente.get("/oportunidades.xlsx"))
    fila = next(f for f in filas if f["Número SICOP"].value == "00000000001")
    for celda, valor in ((fila["Número SICOP"], "00000000001"), (fila["N.º de procedimiento"], "0002026LD-000214-0000900001")):
        assert celda.data_type == "s" and celda.number_format == "@" and celda.value == valor
    # Sin dato: celda vacía, pero con formato texto (lo que se escriba después tampoco se vuelve número).
    vacia = next(f["N.º de procedimiento"] for f in filas if f["Número SICOP"].value == "00000000002")
    assert vacia.value is None and vacia.number_format == "@"


def test_respeta_los_filtros_activos(cliente):
    _, _, alta = _hoja(cliente.get("/oportunidades.xlsx?prioridad=alta"))
    assert sorted(f["Número SICOP"].value for f in alta) == ["00000000001", "00000000004"]
    _, _, busqueda = _hoja(cliente.get("/oportunidades.xlsx?prioridad=todas&q=sillas"))
    assert busqueda == []  # «Compra de sillas» no coincide con ninguna familia: no está en «Todas» las prioridades
    libro, _, buscadas = _hoja(cliente.get("/oportunidades.xlsx?prioridad=todas&q=flotilla"))
    assert [f["Número SICOP"].value for f in buscadas] == ["00000000002"]
    filtros = {fila[0].value: fila[1].value for fila in libro["Filtros"].iter_rows()}
    assert filtros["Prioridad"] == "Todas" and filtros["Búsqueda"] == "flotilla" and filtros["Oportunidades"] == "1"


def test_el_boton_lleva_los_filtros_de_la_pantalla(cliente):
    html = cliente.get("/oportunidades?prioridad=alta&q=gps").text
    assert 'href="/oportunidades.xlsx?prioridad=alta&amp;q=gps"' in html and "Exportar a Excel" in html


def test_ningun_texto_se_guarda_como_formula(cliente, db):
    with closing(connect(db)) as conn:
        conn.execute("UPDATE tenders SET descripcion = '=HYPERLINK(\"http://malo\") monitoreo GPS de vehículos'"
                     " WHERE nro_sicop = '00000000001'")
        conn.commit()
    _, _, filas = _hoja(cliente.get("/oportunidades.xlsx?q=00000000001"))
    objeto = filas[0]["Objeto del cartel"]
    assert objeto.data_type == "s" and objeto.value.startswith("=")


def test_montos_por_moneda_sin_mezclar(cliente, db):
    with closing(connect(db)) as conn:
        conn.executemany(
            "INSERT INTO tender_lines (nro_sicop, numero_linea, estimado_linea, moneda, content_hash, first_seen, last_seen)"
            " VALUES ('00000000001', ?, ?, ?, 'h', 't', 't')", [("1", "1000.50", "CRC"), ("2", "250", "USD"), ("3", "10", "CRC")])
        conn.commit()
    _, titulos, filas = _hoja(cliente.get("/oportunidades.xlsx"))
    assert "Monto estimado CRC" in titulos and "Monto estimado USD" in titulos
    fila = next(f for f in filas if f["Número SICOP"].value == "00000000001")
    assert Decimal(str(fila["Monto estimado CRC"].value)) == Decimal("1010.50")
    assert fila["Monto estimado USD"].value == 250


def test_exportar_exige_sesion(db, monkeypatch):
    monkeypatch.setenv("DATABASE_PATH", str(db))
    respuesta = TestClient(app).get("/oportunidades.xlsx", follow_redirects=False)
    assert respuesta.status_code in (302, 303, 307) and "/login" in respuesta.headers["location"]
